"""build_docx.py: the manuscript and the Supplementary Material as .docx, from the same Markdown the
LaTeX build uses.

Ecological Informatics has no Word template of its own (Elsevier's "Your Paper Your Way"), so the
documents follow Elsevier's general manuscript conventions: A4, 2.5 cm margins, Times New Roman 12 pt,
double spacing, continuous line numbers and page numbers. The manuscript carries the title page,
abstract, keywords, sections 1-6, the Declarations (07_declarations.md), references (Elsevier Harvard,
via pandoc citeproc) and the figure captions. Figures, the graphical abstract and the highlights are
separate files and are not embedded.

Drafting notes (Markdown blockquotes) and HTML comments are dropped, as in build_tex.mjs. Display
equations are numbered in order of appearance and "[#eq:x]" becomes "Eq. (n)"; the supplement's
references to a manuscript equation read "Eq. (n) of the main text".

Usage (needs pandoc; pypandoc_binary and python-docx supply it):
    python paper/tex/build_docx.py
        -> paper/submission/manuscript.docx, paper/submission/supplementary_material.docx
"""
import json, re, subprocess
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = Path(__file__).resolve().parent
P = HERE.parent
OUT = P / "submission"
OUT.mkdir(exist_ok=True)
FM = json.loads((P / "frontmatter.json").read_text(encoding="utf-8"))
BODY = ["01_introduction", "02_related_work", "03_methods", "04_results", "05_discussion",
        "06_conclusions", "07_declarations"]
A4_W, A4_H, MARGIN = 11906, 16838, 1418            # twips; 1418 twips = 2.5 cm
TEXT_WIDTH = A4_W - 2 * MARGIN


def clean(md):
    md = re.sub(r"<!--[\s\S]*?-->", "", md)
    lines = [l for l in md.splitlines() if not re.match(r"^\s*>", l)]
    lines = [l for l in lines if not re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", l)]
    return "\n".join(lines)


def equations(md, eq, external=None):
    """Number labelled display equations into `eq`; resolve [#eq:x] against `eq`, then `external`."""
    def block(m):
        eq[m.group(1)] = len(eq) + 1
        return f"$$\n{m.group(2).strip()} \\qquad ({eq[m.group(1)]})\n$$"
    md = re.sub(r"```math\s*\{#(eq:[\w-]+)\}\s*\n([\s\S]*?)\n```", block, md)
    md = re.sub(r"```math\s*\n([\s\S]*?)\n```", lambda m: f"$$\n{m.group(1).strip()}\n$$", md)

    def ref(m):
        l = m.group(1)
        if l in eq:
            return f"Eq. ({eq[l]})"
        assert external and l in external, f"equation reference without an equation: {l}"
        return f"Eq. ({external[l]}) of the main text"
    md = re.sub(r"\[#(eq:[\w-]+)\]", ref, md)
    return re.sub(r"\$`([^`]+)`\$", r"$\1$", md)                   # GitLab inline math


# ---------------------------------------------------------------- manuscript
EQ = {}
body = equations("\n\n".join(clean((P / f"{f}.md").read_text(encoding="utf-8")) for f in BODY), EQ)
body = re.sub(r"(?m)^# (\d+)\. ", r"# \1 ", body)                # keep the section numbers as text
abstract = re.sub(r"(?m)^# Abstract\s*$", "", clean((P / "00_abstract.md").read_text(encoding="utf-8"))).strip()
title_page = f"""---
title: "{FM['title']}"
lang: en-GB
---

Emrehan Metin^a^, Yunus Emre Cogurcu^a,\\*^

^a^ Çukurova University, Adana, Türkiye

^\\*^ Corresponding author. E-mail: ycogurcu@cu.edu.tr

@@WORDCOUNT@@

# Abstract

{abstract}

**Keywords:** {'; '.join(FM['keywords'])}
"""


def captions():
    tex = "\n".join(l for l in (P / "figure_captions.tex").read_text(encoding="utf-8").splitlines()
                    if not l.lstrip().startswith("%"))
    out, i = [], 0
    while (i := tex.find("\\caption{", i)) >= 0:
        j, depth = i + len("\\caption{"), 1
        while depth:
            depth += {"{": 1, "}": -1}.get(tex[j], 0)
            j += 1
        out.append(tex[i + len("\\caption{"):j - 1])
        i = j
    labels = re.findall(r"\\label\{(fig:[^}]+)\}", tex)

    def refs(c):
        c = re.sub(r"\\ref\{fig:([^}]+)\}", lambda m: str(labels.index("fig:" + m.group(1)) + 1), c)
        return re.sub(r"\\ref\{(?:sec|tab):([^}]+)\}", r"\1", c)
    return [pypandoc.convert_text(refs(c), "markdown", format="latex", extra_args=["--wrap=none"]).strip()
            for c in out]


caps = captions()
fig_md = "# Figure captions\n\n" + "\n\n".join(f"**Fig. {n}.** {c}" for n, c in enumerate(caps, 1))
ms_md = title_page + "\n\n" + body + "\n\n# References\n\n::: {#refs}\n:::\n\n" + fig_md + "\n"

# ---------------------------------------------------------------- supplement
SEQ = {}
sup = clean((P / "supplementary.md").read_text(encoding="utf-8"))
sup = equations(sup, SEQ, external=EQ)
sup = re.sub(r"\A# Supplementary Material\s*\n", "", sup)
sup_md = (f'---\ntitle: "Supplementary Material"\nsubtitle: "{FM["title"]}"\nlang: en-GB\n---\n\n' + sup
          + "\n\n# References\n\n::: {#refs}\n:::\n")


# ---------------------------------------------------------------- Word
def reference_doc(path):
    path.write_bytes(subprocess.run([pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
                                    capture_output=True, check=True).stdout)
    d = Document(path)
    for s in d.styles:
        try:
            f = s.font
        except AttributeError:
            continue
        if f is None:
            continue
        f.name = "Times New Roman"
        rpr = s.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Times New Roman")
    names = [s.name for s in d.styles]
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography", "Abstract"):
        if name in names:
            d.styles[name].font.size = Pt(12)
            d.styles[name].paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for name, size in (("Title", 16), ("Subtitle", 13), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        if name in names:
            d.styles[name].font.size = Pt(size)
            d.styles[name].font.bold = True
            d.styles[name].font.color.rgb = None
    sect = d.sections[0]
    sect.page_width, sect.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sect, side, Cm(2.5))
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
    sect._sectPr.append(ln)
    fp = sect.footer.paragraphs[0] if sect.footer.paragraphs else sect.footer.add_paragraph()
    fp.alignment = 1
    r = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._r.append(e)
    d.save(path)


def to_docx(md, out):
    ref = OUT / "_reference.docx"
    reference_doc(ref)
    pypandoc.convert_text(
        md, "docx", format="markdown+tex_math_dollars+pipe_tables+footnotes+superscript",
        outputfile=str(out),
        extra_args=["--citeproc", f"--bibliography={P / 'REFERENCES.bib'}", f"--csl={HERE / 'elsevier-harvard.csl'}",
                    f"--reference-doc={ref}", "--resource-path=" + str(P), "--columns=1000"])
    ref.unlink()
    # pandoc writes an equal-width grid; size each column by its longest cell, so no renderer wraps a
    # short interval such as [+0.060, +0.073] across two lines
    d = Document(out)
    for t in d.tables:
        need = [max(max(len(c.text.strip()) for c in col.cells), 4) for col in t.columns]
        widths = [round(TEXT_WIDTH * n / sum(need)) for n in need]
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                tcpr = c._tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
                tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(w))
    d.save(out)
    return Document(out)


def main_text_words(doc):
    """Words of Sections 1-6 as Word holds them: from the '1 Introduction' heading to the Declarations
    heading, table cells excluded (tables are not paragraphs), figure captions, references and
    declarations outside the range. Table captions and table notes, which are paragraphs, are counted."""
    n, on = 0, False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading 1"):
            if p.text.startswith("1 "):
                on = True
            elif p.text == "Declarations":
                break
        if on:
            n += len(p.text.split())
    return n


ms_doc = to_docx(ms_md.replace("@@WORDCOUNT@@\n", ""), OUT / "manuscript.docx")
WORDS = main_text_words(ms_doc)
WORDLINE = (f"**Word count:** {WORDS:,} words in the main text, excluding references, tables, "
            "figure captions and declarations.")
ms_md = ms_md.replace("@@WORDCOUNT@@", WORDLINE)
(OUT / "manuscript_assembled.md").write_text(ms_md, encoding="utf-8")
ms_doc = to_docx(ms_md, OUT / "manuscript.docx")
assert main_text_words(ms_doc) == WORDS
sup_doc = to_docx(sup_md, OUT / "supplementary_material.docx")
# the highlights are a separate, editable submission file
(OUT / "highlights.tex").write_bytes((P / "highlights.tex").read_bytes())

for name, doc, eqs in (("manuscript.docx", ms_doc, EQ), ("supplementary_material.docx", sup_doc, SEQ)):
    print(f"{name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(eqs)} numbered equations")
print(WORDLINE.replace("**", ""))
print(f"{len(caps)} figure captions")
