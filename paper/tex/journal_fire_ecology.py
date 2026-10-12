"""journal_fire_ecology.py: the Fire Ecology (SpringerOpen) build profile, `build_docx.py --journal fire-ecology`.

One source, several profiles: the section Markdown (01, 03-07), figure_captions.tex, supplementary.md, the
figures and REFERENCES.bib are the ones the Natural Hazards build reads. Fire Ecology's own items are the
structured abstract (00_abstract_fire_ecology.md) and the title and keywords under "fire_ecology" in
frontmatter.json; everything else that differs lives here, as a transformation of the shared sources.

Fire Ecology's requirements, read 2026-10-12 from the journal's submission guidelines, the Research
article page and the "Fire Ecology Author Guidelines" document linked from them:
- Research article of about 8,000 words, figures, tables and references excluded; a manuscript more than
  10 % longer needs the Editor-in-Chief's agreement before submission;
- Word file, double spaced, left-justified with a ragged right margin, continuous line and page numbers,
  no page breaks; figures and tables embedded after their first reference;
- title page: title in bold sentence case, at most 10 words, no numbers, abbreviations or acronyms;
  authors with superscript affiliation numbers, the corresponding author marked by an asterisk directly
  after the number; ORCID iDs; affiliations; the corresponding author's e-mail;
- "Abstract" in bold with the bold subheadings Background, Results and Conclusions, at most 500 words,
  no citations; "Keywords", at most 10, in alphabetical order, none already in the title;
- sections Background, Methods (active voice), Results (past tense), Discussion, Conclusions; then
  Abbreviations, Acknowledgements (with the use of AI), Author contributions (CRediT), Data availability
  and Declarations (ethics approval and consent to participate, consent for publication, competing
  interests, funding); references in the basic Springer style, alphabetical;
- figure titles of at most 15 words and legends of at most 300; table titles of at most 15 words above
  the table and legends of at most 300 below it; no colour or shading in tables and no commas in
  numbers in tables; figures uploaded as files (PDF accepted, at most 10 MB each);
- supplementary material as "Additional file 1", cited in the text and listed with its file name,
  format, title and description;
- a cover letter with the article type and title, the fit to the scope, novelty and management
  relevance, competing interests, approval by all authors, the statement that the work is not published
  or under consideration elsewhere, every journal that rejected it before, and at least five suggested
  reviewers with affiliation and e-mail.

Output, paper/submission_fire_ecology/:
    manuscript.docx          figures and tables embedded, references in the basic Springer style
    Additional_file_1.pdf    the Supplementary Material with its own reference list
    Fig1.pdf ... Fig8.pdf    the eight vector figures, as drawn and checked by paper/figures/
    cover_letter_draft.docx  unsigned draft, reviewers left out (they are not public)
    MANIFEST.md              file, size, SHA-256, source commit
  local only (gitignored): suggested_reviewers.md, and cover_letter.docx built from the draft with the
  reviewers from that file, for signing.

Checks run at build time (any failure stops the build): the title, abstract and keyword rules above;
the main-text word count (fails above 8,800, the 10 % tolerance); figure and table titles and legends;
every figure and table cited in body prose, first cited in numerical order and placed after its first
citation; every citation resolved; every abbreviation defined in the text listed in Abbreviations and
every listed one used; Additional file 1 cited; decimals: every decimal of the sources reaches the
assembled Markdown, and the Markdown and the .docx hold the same decimals.
"""
import datetime
import hashlib
import json
import re
import shutil
import subprocess
from collections import Counter
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import build_docx as B
import journal_fire as F

HERE = Path(__file__).resolve().parent
P = B.P
ROOT = P.parent
OUT = P / "submission_fire_ecology"
CSL = B.CSL                                         # springer-basic-author-date.csl
FIGS = P / "figures"
FE = B.FM["fire_ecology"]
AUTHORS = [("Emrehan Metin", "EM", "emrehann17@gmail.com", None),
           ("Yunus Emre Cogurcu", "YEC", "ycogurcu@cu.edu.tr", "0000-0002-9229-9657")]
CORRESPONDING = "ycogurcu@cu.edu.tr"
AFFILIATION = ("Department of Computer Engineering, Faculty of Engineering, Çukurova University, "
               "01330 Sarıçam, Adana, Türkiye")
WORD_TARGET, WORD_LIMIT = 8000, 8800                # about 8,000; more than 10 % longer needs the EiC
ABBREVIATIONS = {
    "AUC": "Area under the curve",
    "CI": "Confidence interval",
    "CORAL": "Correlation alignment",
    "ESA": "European Space Agency",
    "LST": "Land surface temperature",
    "MODIS": "Moderate Resolution Imaging Spectroradiometer",
    "NDVI": "Normalised difference vegetation index",
    "ROC": "Receiver operating characteristic",
    "TVDI": "Temperature-Vegetation Dryness Index",
    "VIIRS": "Visible Infrared Imaging Radiometer Suite",
}
ADDITIONAL = ("Sections, tables and figures numbered with S (for example Section S1.23, Table S21 and Fig. S1) "
              "are in Additional file 1.")
# where the first pointer to the supplement goes: after the paragraph that first cites an S section
ADDITIONAL_AFTER = "after the first results; they\nare exploratory, and we report them as such."


# ---------------------------------------------------------------- sources
def words(s):
    return len(re.sub(r"[*_`#|]", " ", s).split())


def check_title(title):
    assert title != B.FM["title"], "Fire Ecology title not set"
    n = len(title.split())
    assert n <= 10, f"title has {n} words; Fire Ecology allows 10"
    assert not re.search(r"\d", title), "title contains a number"
    assert not re.search(r"\b[A-Z]{2,}\b", title), "title contains an acronym"
    assert title[0].isupper() and title[1:] == title[1:].lower(), "title is not in sentence case"


def structured_abstract():
    src = B.clean((P / "00_abstract_fire_ecology.md").read_text(encoding="utf-8"), "abstract (Fire Ecology)")
    text = re.sub(r"(?m)^# Abstract\s*$", "", src).strip()
    paras = [p.strip() for p in text.split("\n\n")]
    heads = [re.match(r"\*\*(\w+)\*\* ", p).group(1) if re.match(r"\*\*(\w+)\*\* ", p) else None for p in paras]
    assert heads == ["Background", "Results", "Conclusions"], f"abstract subheadings {heads}"
    assert "@" not in text, "citation in the abstract"
    n = words(text)
    assert n <= 500, f"abstract has {n} words; Fire Ecology allows 500"
    return text, n


def check_keywords(kw, title):
    assert len(kw) <= 10, f"{len(kw)} keywords; Fire Ecology allows 10"
    assert kw == sorted(kw, key=str.casefold), f"keywords not in alphabetical order: {kw}"
    clash = [k for k in kw if k.casefold() in title.casefold()]
    assert not clash, f"keywords already in the title: {clash}"


def tables(body):
    """'**Table n. Title.** legend' + pipe table + optional table note -> title above (Springer
    'Table n Title'), table, legend and note below; commas dropped from numbers inside the table."""
    out, n = [], 0
    pat = re.compile(r"\*\*Table (\d+)\. (.+?)\.\*\*\s*(.*?)\n\n((?:\|[^\n]*\n?)+)(?:\n\*Table note\.\* (.*?)(?=\n\n|\Z))?", re.S)
    items = []

    def fix(m):
        nonlocal n
        n += 1
        num, title, legend, table, note = m.groups()
        assert int(num) == n, f"table {num} out of order"
        assert words(title) <= 15, f"Table {num} title has {words(title)} words; the limit is 15"
        legend = " ".join(legend.split())
        below = legend + (" " + " ".join(note.split()) if note else "")
        assert words(below) <= 300, f"Table {num} legend has {words(below)} words; the limit is 300"
        table = re.sub(r"(?<=\d),(?=\d{3}\b)", "", table)                  # 2,802 -> 2802 inside tables
        below = re.sub(r"(?<=\d),(?=\d{3}\b)", "", below)
        items.append((int(num), words(title), words(below)))
        return f"**Table {num}** {title}\n\n{table.rstrip()}\n\n{below}"
    body = pat.sub(fix, body)
    assert not re.search(r"\*\*Table \d+\. ", body), "a table caption was not converted"
    return body, items


def figure_captions():
    caps = []
    for n, c in enumerate(B.captions(), 1):
        m = re.match(r"\*\*(.+?)\*\*\s*(.*)", c, re.S)
        assert m, f"Fig. {n} caption does not start with a bold title"
        title, legend = m.group(1).strip(), " ".join(m.group(2).split())
        assert words(title) <= 15, f"Fig. {n} title has {words(title)} words; the limit is 15"
        assert words(legend) <= 300, f"Fig. {n} legend has {words(legend)} words; the limit is 300"
        caps.append((title, legend))
    return caps


def insert_figures(body, caps, embed):
    ends = []
    for n in range(1, len(caps) + 1):
        ends.append((F.first_cite_paragraph_end(body, rf"\bFigs?\. (?:\d+(?:, | and | to ))*{n}\b(?!\.\d)"), n))
    out, last = [], 0
    for end, n in sorted(ends):
        path, width = embed[n]
        title, legend = caps[n - 1]
        out += [body[last:end], f"\n\n![**Fig. {n} {title}** {legend}]({path.as_posix()}){{width={width:.2f}cm}}\n"]
        last = end
    return "".join(out) + body[last:]


def main_text_words(body):
    """Background to Conclusions without tables, table titles, legends and notes, figure captions and
    equations; headings count."""
    b = re.sub(r"!\[.*?\]\([^)]*\)\{[^}]*\}", "", body, flags=re.S)         # figures with captions
    b = re.sub(r"\*\*Table \d+\*\* [^\n]*\n\n(?:\|[^\n]*\n?)+\n?[^\n]*(?:\n[^\n]+)*", "", b)   # title, table, legend
    b = re.sub(r"\$\$.*?\$\$", "", b, flags=re.S)
    return words(b)


def _expand(s):
    """'1', '1, 2 and 4', '6 to 8' -> the numbers they cite."""
    parts = re.split(r"(, | and | to )", s)
    nums, out = [int(x) for x in parts[::2]], set()
    out.update(nums)
    for i, sep in enumerate(parts[1::2]):
        if sep == " to ":
            out.update(range(nums[i], nums[i + 1] + 1))
    return out


def check_display_items(md, n_fig, n_tab):
    """In the assembled Markdown: every item cited in body prose (not in a heading, a table, a figure
    caption or a table title), first cited in numerical order, and placed after that citation."""
    body = md[md.index("\n# 1 Background"):md.index("\n# Abbreviations")]
    blocks = [(m.start(), m.group(0)) for m in re.finditer(r"[^\n]+(?:\n[^\n]+)*", body)]
    report = {}
    for label, cite, place, n_items in (("Fig", r"\bFigs?\. ((?:\d+(?:, | and | to ))*\d+)\b(?!\.\d)", r"!\[\*\*Fig\. (\d+) ", n_fig),
                                        ("Table", r"\bTables? ((?:\d+(?:, | and | to ))*\d+)\b", r"\*\*Table (\d+)\*\* ", n_tab)):
        placed = {int(m.group(1)): m.start() for m in re.finditer(place, body)}
        first = {}
        for pos, t in blocks:
            if t.startswith(("|", "#", "![", "$$")) or re.match(r"\*\*Table \d+\*\* ", t):
                continue
            for m in re.finditer(cite, t):
                for n in _expand(m.group(1)):
                    first.setdefault(n, pos)
        want = set(range(1, n_items + 1))
        assert set(placed) == want, f"{label} placed {sorted(placed)} != 1..{n_items}"
        assert want <= set(first), f"{label}s not cited in body prose: {sorted(want - set(first))}"
        order = [first[n] for n in sorted(want)]
        assert order == sorted(order), f"{label}s first cited out of numerical order"
        late = [n for n in want if placed[n] < first[n]]
        assert not late, f"{label}s placed before their first citation: {late}"
        report[label] = len(want)
    return report


def check_abbreviations(text):
    """Every abbreviation defined in the text, "(LST)" or "(ROC-AUC)", is listed, and every listed one is used."""
    defined = {part for d in re.findall(r"\(([A-Z]{2,}(?:-[A-Z]{2,})*)\)", text) for part in d.split("-")}
    listed = set(ABBREVIATIONS)
    missing = defined - listed
    assert not missing, f"abbreviations defined in the text but not listed: {sorted(missing)}"
    unused = [a for a in listed if not re.search(rf"\b{a}\b", text)]
    assert not unused, f"listed abbreviations not used in the text: {unused}"


def docx_left_aligned(path):
    d = Document(path)
    for s in d.styles:
        pf = getattr(s, "paragraph_format", None)
        if pf is not None and s.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography",
                                        "Abstract", "Image Caption", "Table Caption", "Captioned Figure"):
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for p in d.paragraphs:
        for br in p._p.iter(qn("w:br")):
            assert br.get(qn("w:type")) != "page", "page break in the manuscript"
    d.save(path)
    B.strip_custom_props(path)


def check_numbers(sources, md, path):
    md = re.sub(r"\{[^{}]*(?:custom-style|width=)[^{}]*\}", "", md)
    src, asm = Counter(F.DEC.findall(sources)), Counter(F.DEC.findall(md))
    lost = src - asm
    assert not lost, f"numbers in the sources that the profile dropped: {dict(lost)}"
    d = Document(path)
    docx, refs = Counter(), False
    for p in d.element.body.iter(qn("w:p")):
        t = F._text(p)
        if t.strip() == "References":
            refs = True
        if not refs:
            docx.update(F.DEC.findall(t))
    md_body = md[:md.index("\n# References")]
    asm_body = Counter(F.DEC.findall(md_body))
    missing, extra = asm_body - docx, docx - asm_body
    assert not missing and not extra, f"Markdown vs .docx decimals: missing {dict(missing)}, extra {dict(extra)}"
    return sum(asm_body.values())


# ---------------------------------------------------------------- build
def build(pdf=True):
    OUT.mkdir(exist_ok=True)
    B.OUT = OUT
    title, kw = FE["title"], FE["keywords"]
    check_title(title)
    check_keywords(kw, title)
    abstract, n_abs = structured_abstract()

    raw = {f: B.clean((P / f"{f}.md").read_text(encoding="utf-8"), f) for f in B.BODY}
    EQ = {}
    body = B.equations(B.narrative("\n\n".join(raw[f] for f in B.BODY)), EQ)
    body = re.sub(r"(?m)^# (\d+)\. ", r"# \1 ", body)
    assert body.count(ADDITIONAL_AFTER) == 1, "anchor for the Additional file 1 pointer not found"
    body = body.replace(ADDITIONAL_AFTER, ADDITIONAL_AFTER + " " + ADDITIONAL)
    heads = re.findall(r"(?m)^# \d+ (.+)$", body)
    assert heads == ["Background", "Methods", "Results", "Discussion", "Conclusions"], heads
    body, tab_items = tables(body)

    # figures: vector PDFs for upload, 300 dpi copies embedded after the first citing paragraph
    caps = figure_captions()
    tmp = OUT / "_embed"
    tmp.mkdir(exist_ok=True)
    embed, fig_src = {}, {}
    for n in range(1, len(caps) + 1):
        src = sorted(FIGS.glob(f"fig{n}_*.pdf"))
        assert len(src) == 1, src
        fig_src[n] = src[0]
        shutil.copyfile(src[0], OUT / f"Fig{n}.pdf")
        assert (OUT / f"Fig{n}.pdf").stat().st_size < 10 * 2 ** 20, f"Fig{n}.pdf exceeds 10 MB"
        _, width = F.rasterise(n, 300, tmp / f"Fig{n}.png")
        embed[n] = (tmp / f"Fig{n}.png", min(width, 16.0))
    n_words = main_text_words(body)
    assert n_words <= WORD_LIMIT, f"main text has {n_words} words; above {WORD_LIMIT} the EiC must agree first"
    body = insert_figures(body, caps, embed)

    decl_src = B.clean((P / f"{B.DECL}.md").read_text(encoding="utf-8"), B.DECL)
    decl = F.sections(decl_src)
    assert set(decl) == {"Funding", "Competing interests", "Author contributions", "Data availability",
                         "Code availability", "Ethics approval", "Use of generative AI"}, sorted(decl)
    assert decl["Ethics approval"].startswith("Not applicable."), decl["Ethics approval"][:40]
    contrib = F.flat(decl["Author contributions"])
    assert contrib.startswith("Stated in CRediT terms."), contrib[:40]
    for name, ini, *_ in AUTHORS:
        contrib = contrib.replace(f"**{name}:**", f"{ini}:")
    contrib = contrib.replace("Stated in CRediT terms. ", "")
    ack = B.ACK.split("\n\n", 1)[1].strip() + " " + F.flat(decl["Use of generative AI"])
    sup_src = B.clean((P / "supplementary.md").read_text(encoding="utf-8"), "supplementary")
    n_sup_tab = len(re.findall(r"(?m)^\*\*Table S\d+\. ", sup_src))
    doi = re.search(r"\{ThermalTwinRepo,[^@]*?doi\s*=\s*\{([^}]+)\}", (P / "REFERENCES.bib").read_text(encoding="utf-8"))
    assert doi, "ThermalTwinRepo has no DOI in REFERENCES.bib"
    software = ("Project name: thermal-twin. Project home page: <https://github.com/dryuemco/thermal-twin>. "
                f"Archived version: <https://doi.org/{doi.group(1)}>. Operating system: platform independent. Programming "
                "language: Python 3.12. Other requirements: NumPy 2.4.4, pandas 3.0.2, scikit-learn 1.9.0. "
                "License: MIT. Any restrictions to use by non-academics: none.")

    names = [f"{name}^1{'\\*' if mail == CORRESPONDING else ''}^" for name, _, mail, _ in AUTHORS]
    orcid = "; ".join(f"{name}, {o}" for name, _, _, o in AUTHORS if o)
    abbrev = "; ".join(f"{k}: {v}" for k, v in sorted(ABBREVIATIONS.items()))
    front = f"""---
lang: en-GB
---

**{title}**

{", ".join(names)}

^1^ {AFFILIATION}

^\\*^ Corresponding author: {CORRESPONDING}

ORCID: {orcid}

Article type: Research. Main text (Background to Conclusions, tables, figures and references excluded):
@@WORDS@@ words. Abstract: {n_abs} words.

**Abstract**

{abstract}

**Keywords** {", ".join(kw)}
"""
    back = f"""# Abbreviations

{abbrev}

# Supplementary Information

Additional file 1 (Additional_file_1.pdf; PDF). Supplementary Material: sensitivity analyses and detail
(Sections S1.1 to S1.23), supporting tables (Section S2), protocol detail and the full list of limitations
(Section S3), the target-label recovery curve (Section S4), Fig. S1 and Tables S1 to S{n_sup_tab}, each with
the source files of its values.

# Acknowledgements

{ack}

# Author contributions

{contrib}

# Data availability

The datasets supporting the conclusions of this article are available in the Zenodo repository,
<https://doi.org/{doi.group(1)}> [@ThermalTwinRepo].

{decl["Data availability"]}

{decl["Code availability"]}

{software}

# Declarations

## Ethics approval and consent to participate

{decl["Ethics approval"]}

## Consent for publication

Not applicable.

## Competing interests

{decl["Competing interests"]}

## Funding

{decl["Funding"]}

# References

::: {{#refs}}
:::
"""
    ms_md = front + "\n\n" + body + "\n\n" + back
    check_abbreviations(abstract + "\n" + body)
    assert "Additional file 1" in body, "Additional file 1 not cited in the text"
    items = check_display_items(ms_md.replace("@@WORDS@@", "0"), len(caps), len(tab_items))
    ms_md = ms_md.replace("@@WORDS@@", f"{n_words:,}")
    (HERE / "manuscript_fire_ecology_assembled.md").write_text(ms_md, encoding="utf-8")   # build record

    out = OUT / "manuscript.docx"
    doc = B.to_docx(ms_md, out, title)
    docx_left_aligned(out)
    d = Document(out)
    allt = "\n".join(F._text(p) for p in d.element.body.iter(qn("w:p")))
    assert "???" not in allt and not re.search(r"(?<![\w.])@[A-Za-z]", allt), "unresolved citation in the .docx"
    cp = d.core_properties
    iso = subprocess.run(["git", "log", "-1", "--format=%cI"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    cp.created = cp.modified = datetime.datetime.fromisoformat(iso).astimezone(datetime.timezone.utc).replace(tzinfo=None)
    d.save(out)
    B.strip_custom_props(out)
    shutil.rmtree(tmp)
    sources = "\n".join([abstract, "\n".join(raw.values()), "\n".join(B.captions()), decl_src])
    n_dec = check_numbers(sources, ms_md, out)

    # ---- Additional file 1: the shared Supplementary Material, with its own reference list
    SEQ = {}
    sup = B.equations(B.narrative(sup_src), SEQ, prefix="S", external=EQ)
    sup = re.sub(r"\A# Supplementary Material\s*\n", "", sup)
    sup_md = (f'---\ntitle: "Additional file 1: Supplementary Material"\nsubtitle: "{title}"\nlang: en-GB\n---\n\n'
              + sup + "\n\n# References\n\n::: {#refs}\n:::\n")
    sup_docx = OUT / "_additional_file_1.docx"
    B.to_docx(sup_md, sup_docx, "Additional file 1: " + title)
    F.balance_tables(sup_docx)
    if pdf:
        F.word_pdf(sup_docx, OUT / "Additional_file_1.pdf")
        assert (OUT / "Additional_file_1.pdf").stat().st_size < 20 * 2 ** 20, "Additional file 1 exceeds 20 MB"
    sup_docx.unlink()

    # ---- cover letter: the committed draft has no reviewers; the local copy gets them
    letter = (HERE / "fire_ecology_cover_letter.md").read_text(encoding="utf-8")
    letter = letter.replace("{{TITLE}}", title).replace("{{WORDS}}", f"{n_words:,}")
    for s in ("has not been published, and it is not under consideration for publication elsewhere",
              "All authors have approved the manuscript for submission",
              "MDPI; manuscript fire-4642098", "declined it without peer review", "Research article",
              "competing interests", "special collection"):
        assert s in F.flat(letter), f"cover letter lacks: {s}"
    stray = set(F.DEC.findall(letter)) - set(F.DEC.findall(abstract))
    assert not stray, f"cover letter numbers not in the abstract: {stray}"
    draft = letter.replace("{{REVIEWERS}}", "\\[Five suggested reviewers with affiliation and e-mail: inserted "
                                             "from the local file suggested_reviewers.md, which is not public.\\]")
    B.to_docx(draft, OUT / "cover_letter_draft.docx", "Cover letter")
    F.letter_format(OUT / "cover_letter_draft.docx")
    rev = OUT / "suggested_reviewers.md"
    if rev.exists():
        block = rev.read_text(encoding="utf-8").split("<!-- cover letter -->", 1)
        assert len(block) == 2, "suggested_reviewers.md lacks the '<!-- cover letter -->' marker"
        entries = re.findall(r"(?m)^\d+\. ", block[1])
        assert len(entries) >= 5, f"{len(entries)} suggested reviewers; Fire Ecology asks for at least five"
        B.to_docx(letter.replace("{{REVIEWERS}}", block[1].strip()), OUT / "cover_letter.docx", "Cover letter")
        F.letter_format(OUT / "cover_letter.docx")

    manifest(fig_src, n_abs, n_words, n_dec, tab_items, items)
    print(f"manuscript.docx: title {len(title.split())} words, abstract {n_abs} words, {len(kw)} keywords, "
          f"main text {n_words} words (target {WORD_TARGET}, limit {WORD_LIMIT}), {len(caps)} figures and "
          f"{len(tab_items)} tables embedded, {n_dec} decimals matched")


def manifest(fig_src, n_abs, n_words, n_dec, tab_items, items):
    commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    dirty = subprocess.run(["git", "status", "--porcelain", "--", "paper", ":!paper/submission_fire_ecology",
                            ":!paper/tex/manuscript_fire_ecology_assembled.md"],
                           cwd=ROOT, capture_output=True, text=True).stdout.strip()
    state = commit + (" (with uncommitted source changes)" if dirty else "")
    src = {"manuscript.docx": "paper/01, 03-07 *.md, 00_abstract_fire_ecology.md, figure_captions.tex, frontmatter.json, REFERENCES.bib",
           "Additional_file_1.pdf": "paper/supplementary.md, Word PDF export",
           "cover_letter_draft.docx": "paper/tex/fire_ecology_cover_letter.md (unsigned, without the reviewers)"}
    for n, f in fig_src.items():
        src[f"Fig{n}.pdf"] = f"{f.relative_to(ROOT).as_posix()} (vector, unchanged)"
    rows = []
    for name in ["manuscript.docx", "Additional_file_1.pdf", *[f"Fig{n}.pdf" for n in fig_src], "cover_letter_draft.docx"]:
        f = OUT / name
        if f.exists():
            rows.append(f"| `{name}` | {f.stat().st_size} | `{hashlib.sha256(f.read_bytes()).hexdigest()}` | {src[name]} |")
    tabs = "; ".join(f"Table {n}: title {t} words, legend {l} words" for n, t, l in tab_items)
    text = f"""# Submission package: Fire Ecology (SpringerOpen)

Research article. Built from commit `{state}` by `python paper/tex/build_docx.py --journal fire-ecology`;
the profile is `paper/tex/journal_fire_ecology.py`. The shared sources were restructured for this journal
on 2026-10-12 (Background merged from the former Introduction and Related work, Methods in the active
voice, about 900 words shorter); the earlier Fire (MDPI) package was built from commit `ab058c2`.

Build-time checks: title {len(FE['title'].split())} words, no numbers or acronyms; structured abstract
{n_abs} words (limit 500); {len(FE['keywords'])} keywords, alphabetical, none in the title; main text
{n_words} words (about {WORD_TARGET} asked, {WORD_LIMIT} at most without the Editor-in-Chief); figure and
table titles at most 15 words and legends at most 300 ({tabs}); {items.get('Fig')} figures and
{items.get('Table')} tables cited in body prose, in order, placed after their first citation; every
citation resolved; abbreviations listed; {n_dec} decimals identical between the assembled Markdown and the
.docx. Run separately: `paper/figures/check_all.py`, `paper/code/check_stale_values.py`,
`paper/code/verify_references.py`.

Open items for the corresponding author: Emrehan Metin's approval of this version (title and text
changed); the cover letter's date and signature (`cover_letter.docx`, local); at least five suggested
reviewers (`suggested_reviewers.md`, local); the article processing charge is paid from the BAP project
after acceptance.

| File | Size (bytes) | SHA-256 | Source |
|---|---:|---|---|
""" + "\n".join(rows) + "\n"
    (OUT / "MANIFEST.md").write_text(text, encoding="utf-8")
