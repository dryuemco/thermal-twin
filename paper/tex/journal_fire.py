"""journal_fire.py: the Fire (MDPI) build profile, run as `build_docx.py --journal fire`.

One source, two profiles: the manuscript Markdown (00-07), figure_captions.tex, supplementary.md,
the figures and REFERENCES.bib are the ones the Natural Hazards build reads. Everything that differs
between the journals lives here, as a transformation of those sources, never as a second copy of
the text. Where the profile rewords a declaration into MDPI's fixed form, it parses the source
sentence and fails if the source has been reworded, so the two builds cannot drift apart silently.

Fire's Instructions for Authors and its Word template (both read 2026-09-30) ask for:
- the MDPI Word template; figures inserted in the text after the paragraph of their first citation,
  captions below figures and above tables, cited as "Figure 1", "Table 1"; figures also uploaded as
  PNG, JPEG or TIFF, preferably at 600 dpi or more;
- front matter: article type, title, authors, affiliations, correspondence, an abstract of about 200
  words at most, 3 to 10 keywords; Highlights are optional and are not used;
- sections Introduction, Materials and Methods, Results, Discussion, Conclusions, numbered "3.1.";
- back matter, in this order: Supplementary Materials, Author Contributions (CRediT, initials),
  Funding, Institutional Review Board Statement, Informed Consent Statement, Data Availability
  Statement, Acknowledgments (with the GenAI statement), Conflicts of Interest, References;
- references numbered in order of first appearance (captions included), [1], [1-3], [1,3], placed
  before punctuation, ACS-like MDPI style (paper/tex/mdpi.csl, Zotero style repository, CC BY-SA);
  citations in the supplement must also appear in the main reference list;
- a graphical abstract (optional) of at least 1100 x 560 px, PNG, JPEG or TIFF.

The template is MDPI's, distributed for submissions only; it is not in the repository. Pass the
downloaded file with --template (a .dot that is really a .dotx zip is accepted and converted).

Output, paper/submission_fire/:
    manuscript.docx        MDPI template, figures embedded (300 dpi copies)
    supplement.pdf         the Supplementary Material, citations numbered as in the main list
    Figure1.png ...        the eight figures at 600 dpi, rasterised from the checked PDFs
    graphical_abstract.png from paper/figures/graphical_abstract.py
    MANIFEST.md            file, size, SHA-256, source commit

Checks run at build time (any failure stops the build):
    abstract <= 200 words; 3 to 10 keywords; no bare narrative citation; every cited key in the
    bibliography; every supplement key cited in the main text;
    citations: every [n] in the manuscript resolves to a reference, every reference is cited, and
    numbers appear first in the order 1, 2, 3, ...; the supplement uses only numbers of the list;
    display items (check 2b of the former verify_tex): every figure and table is cited in body prose
    (not only in a caption, a table cell or the back matter), first cited in numerical order, and
    placed after its first citation;
    numbers: every decimal in the sources survives into the assembled Markdown, and the Markdown and
    the .docx hold the same decimals.
"""
import hashlib
import json
import re
import shutil
import subprocess
import zipfile
from collections import Counter
from pathlib import Path

import pypandoc
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import build_docx as B

HERE = Path(__file__).resolve().parent
P = B.P
ROOT = P.parent
OUT = P / "submission_fire"
CSL = HERE / "mdpi.csl"
FIGS = P / "figures"
GA = FIGS / "graphical_abstract.png"
INDENT, TEXT_W, FULL_W = 2608, 7857, 10466       # twips: template body indent, body width, full width
SECTION_RENAME = {"Data and methods": "Materials and Methods"}
# Natural Hazards collects the statements under "Statements and Declarations"; Fire has no such section
DECLARATIONS_REF = {"01_introduction": [("re-run\n(Declarations).", "re-run\n(Data Availability Statement).")],
                    "supplementary": [("given\nin the Declarations.",
                                       "given\nin the Data Availability Statement of the main article.")]}
AUTHORS = [("Emrehan Metin", "E.M."), ("Yunus Emre Cogurcu", "Y.E.C.")]
CORRESPONDING = "ycogurcu@cu.edu.tr"
AFFILIATION = ("Department of Computer Engineering, Faculty of Engineering, Çukurova University, "
               "01330 Sarıçam, Adana, Türkiye")
FUNDER_REGISTRY = "Çukurova University"   # Crossref Funder Registry 501100002964, alt-name of "Çukurova Üniversitesi"
CREDIT = ["Conceptualization", "Methodology", "Software", "Validation", "Formal analysis", "Investigation",
          "Resources", "Data curation", "Writing (original draft)", "Writing (review and editing)",
          "Visualization", "Supervision", "Project administration", "Funding acquisition"]
CREDIT_MDPI = {"Writing (original draft)": "writing—original draft preparation",
               "Writing (review and editing)": "writing—review and editing"}
DEC = re.compile(r"\d+\.\d+")
M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"


# ---------------------------------------------------------------- sources
def figures_to_words(md):
    md = re.sub(r"\bFigs\. ", "Figures ", md)
    return re.sub(r"\bFig\. ", "Figure ", md)


def declarations_refs(md, name):
    for old, new in DECLARATIONS_REF.get(name, []):
        assert md.count(old) == 1, f"{name}: {old!r} not found once"
        md = md.replace(old, new)
    assert "Declarations" not in md, f"{name}: a reference to the Declarations is left"
    return md


# MDPI prints abbreviated journal names (ISO 4, as MDPI writes them); the shared bibliography keeps the
# full names that the Springer style needs, so the abbreviations are added to Fire's copy only
ISO4 = {
    "AIMS Environmental Science": "AIMS Environ. Sci.", "Discover Artificial Intelligence": "Discov. Artif. Intell.",
    "Earth Science Informatics": "Earth Sci. Inform.", "Earth System Science Data": "Earth Syst. Sci. Data",
    "Ecography": "Ecography", "Ecological Informatics": "Ecol. Inform.", "Ecological Modelling": "Ecol. Model.",
    "Ecological Monographs": "Ecol. Monogr.", "Ecology": "Ecology", "Environmental Modelling \\& Software":
    "Environ. Model. Softw.", "Environmental Reviews": "Environ. Rev.", "Evolution": "Evolution", "Fire": "Fire",
    "Fire Ecology": "Fire Ecol.", "Forest Ecology and Management": "For. Ecol. Manag.",
    "Frontiers in Ecology and the Environment": "Front. Ecol. Environ.", "Geomatics": "Geomatics",
    "Global Ecology and Biogeography": "Glob. Ecol. Biogeogr.",
    "IEEE Geoscience and Remote Sensing Magazine": "IEEE Geosci. Remote Sens. Mag.",
    "IEEE Transactions on Geoscience and Remote Sensing": "IEEE Trans. Geosci. Remote Sens.",
    "International Journal of Wildland Fire": "Int. J. Wildland Fire",
    "ISPRS Journal of Photogrammetry and Remote Sensing": "ISPRS J. Photogramm. Remote Sens.",
    "Machine Learning": "Mach. Learn.", "Methods in Ecology and Evolution": "Methods Ecol. Evol.",
    "Natural Hazards and Earth System Sciences": "Nat. Hazards Earth Syst. Sci.",
    "Nature Communications": "Nat. Commun.", "Pattern Recognition": "Pattern Recognit.",
    "Remote Sensing": "Remote Sens.", "Remote Sensing of Environment": "Remote Sens. Environ.",
    "Trends in Ecology \\& Evolution": "Trends Ecol. Evol.",
}
NOT_JOURNALS = {"Zenodo", "NASA Land Processes Distributed Active Archive Center", "arXiv preprint",
                "U.S. Geological Survey", "European Space Agency"}


def fire_bib(dest):
    """Fire's copy of REFERENCES.bib: ISO 4 journal abbreviations added, and the "Proceedings of the"
    prefix dropped from the one booktitle, because MDPI's style prints "In Proceedings of the" itself."""
    bib = (P / "REFERENCES.bib").read_text(encoding="utf-8")
    bib, n = re.subn(r"(booktitle\s*=\s*\{)Proceedings of the ", r"\1", bib)
    assert n == 1, n

    def short(m):
        name = " ".join(re.sub(r"[{}]", "", m.group(2)).split())
        if name in NOT_JOURNALS:
            return m.group(0)
        assert name in ISO4, f"journal without an ISO 4 abbreviation in journal_fire.ISO4: {name}"
        return f"{m.group(0)},\n  shortjournal = {{{ISO4[name]}}}"
    bib, n = re.subn(r"(journal\s*=\s*)\{((?:[^{}]|\{[^{}]*\})*)\}", short, bib)
    assert n > 0
    dest.write_text(bib, encoding="utf-8")
    return dest


def headings(md):
    for old, new in SECTION_RENAME.items():
        md, n = re.subn(rf"(?m)^(# \d+\. ){re.escape(old)}\s*$", rf"\g<1>{new}", md)
        assert n == 1, f"section heading '{old}' not found"
    md = re.sub(r"(?m)^(#{2,3}) (\d+(?:\.\d+)+) ", r"\1 \2. ", md)          # 3.1 -> 3.1.
    return md


def sections(md):
    """'## Name' -> body text, for 07_declarations.md."""
    parts = re.split(r"(?m)^## (.+?)\s*$", md)
    return {parts[i].strip(): parts[i + 1].strip() for i in range(1, len(parts), 2)}


def flat(s):
    return " ".join(s.split())


def author_contributions(text):
    text = flat(text)
    assert text.startswith("Stated in CRediT terms."), text[:60]
    assert text.endswith("Both authors read and approved the final manuscript."), text[-60:]
    roles = {}
    for name, initials in AUTHORS:
        m = re.search(rf"\*\*{re.escape(name)}:\*\* (.+?)\.(?= \*\*| Both authors)", text)
        assert m, f"no CRediT roles for {name}"
        for r in (x.strip() for x in m.group(1).split(", ")):
            assert r in CREDIT, f"unknown CRediT role '{r}'"
            roles.setdefault(r, []).append(initials)
    items = []
    for r in CREDIT:
        if r in roles:
            who = roles[r][0] if len(roles[r]) == 1 else ", ".join(roles[r][:-1]) + " and " + roles[r][-1]
            items.append(f"{CREDIT_MDPI.get(r, r.lower())}, {who}")
    s = "; ".join(items)
    return (s[0].upper() + s[1:] + ("" if s.endswith(".") else ".")      # initials already end in "."
            + " All authors have read and agreed to the published version of the manuscript.")


def funding_and_coi(funding, competing):
    f = flat(funding)
    m = re.fullmatch(r"This work was supported by the ((.+?) Scientific Research Projects Coordination Unit) "
                     r"(\(.+?\)) under its (.+?) scheme, project code (FKB-\d{4}-\d+) (\(.+?\))\. "
                     r"(The funder had no role .+\.)", f)
    assert m, "Funding paragraph reworded; update journal_fire.funding_and_coi"
    unit, uni, native, scheme, code, title, role = m.groups()
    assert uni == FUNDER_REGISTRY, uni
    fund = f"This research was funded by the {unit} {native} under its {scheme} scheme, grant number {code} {title}."
    assert flat(competing) == "The authors have no relevant financial or non-financial interests to disclose.", competing
    return fund, "The authors declare no conflicts of interest. " + role


def genai(text, methods):
    t = flat(text)
    m = re.fullmatch(r"Claude \(Anthropic\) was used in the research and in preparing this manuscript, as described "
                     r"in Section 3\.13\. It was used to (.+?)\. The authors reviewed and edited all of this "
                     r"material and take full responsibility for the content of the publication\.", t)
    assert m, "Use of generative AI paragraph reworded; update journal_fire.genai"
    v = re.search(r"Claude \((Anthropic; [^)]+)\)", flat(methods))
    assert v, "tool and version not found in Section 3.13"
    return (f"During the preparation of this study and manuscript, the authors used Claude ({v.group(1)}) for the "
            f"purposes described in Section 3.13: to {m.group(1)}. The authors have reviewed and edited the output "
            "and take full responsibility for the content of this publication.")


def supplementary_items(sup):
    items = []
    for kind, pat in (("Figure", r"!\[\*\*Fig\. (S\d+)\. (.+?)\*\*"), ("Table", r"\*\*Table (S\d+)\. (.+?)\*\*")):
        for n, title in re.findall(pat, sup, re.S):
            title = flat(re.sub(r"[`*]", "", title)).rstrip(".")
            items.append((kind, int(n[1:]), f"{kind} {n}: {title}"))
    items.sort()
    tables = [n for k, n, _ in items if k == "Table"]
    assert tables == list(range(1, len(tables) + 1)), f"supplementary tables not S1..S{len(tables)}: {tables}"
    return "; ".join(t for _, _, t in items)


def cite_order(md, bib):
    """Reference keys in citeproc's numbering order, from pandoc's JSON; fails on an unresolved key."""
    r = subprocess.run([pypandoc.get_pandoc_path(), "-f", "markdown+tex_math_dollars+pipe_tables+footnotes+superscript",
                        "-t", "json", "--citeproc", f"--bibliography={bib}", f"--csl={CSL}"],
                       input=md.encode("utf-8"), capture_output=True, check=True)
    err = r.stderr.decode("utf-8", "replace")
    assert "not found" not in err, err
    return re.findall(r'"ref-([^"]+)"', r.stdout.decode("utf-8"))


def number_citations(md, order):
    """[@a; @b] -> [n1,n2] with the main list's numbers, ranges of three or more collapsed (MDPI)."""
    num = {k: i + 1 for i, k in enumerate(order)}

    def group(m):
        keys = re.findall(r"@([A-Za-z][\w]*)", m.group(0))
        assert re.fullmatch(r"\[\s*@[\w]+(?:\s*;\s*@[\w]+)*\s*\]", m.group(0)), f"citation with a locator: {m.group(0)}"
        missing = [k for k in keys if k not in num]
        assert not missing, f"supplement cites keys absent from the main reference list: {missing}"
        ns, out, i = sorted(num[k] for k in keys), [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            out += [f"{ns[i]}–{ns[j]}"] if j - i >= 2 else [str(x) for x in ns[i:j + 1]]
            i = j + 1
        return "[" + ",".join(out) + "]"
    return re.sub(r"\[[^\[\]]*@[^\[\]]*\]", group, md)


# ---------------------------------------------------------------- figures
def rasterise(n, dpi, out):
    import fitz
    src = sorted(FIGS.glob(f"fig{n}_*.pdf"))
    assert len(src) == 1, f"figure {n}: expected one fig{n}_*.pdf, found {src}"
    doc = fitz.open(src[0])
    page = doc[0]
    page.get_pixmap(dpi=dpi, alpha=False, colorspace=fitz.csRGB).save(out)
    width_cm = page.rect.width / 72 * 2.54
    doc.close()
    return src[0], width_cm


def first_cite_paragraph_end(body, pat):
    """End offset of the prose paragraph that first matches `pat`; a table that follows is skipped."""
    for m in re.finditer(pat, body):
        start = body.rfind("\n\n", 0, m.start()) + 2
        block = body[start:m.start()]
        if block.lstrip().startswith(("|", "#", "![")):
            continue                                              # table cell, heading or figure: not prose
        end = body.find("\n\n", m.end())
        end = len(body) if end < 0 else end
        rest = body[end:].lstrip("\n")
        if rest.startswith("|"):                                  # the paragraph captions a table: after it
            tab_end = re.search(r"\n(?!\|)", rest)
            end = len(body) - len(rest) + (tab_end.start() if tab_end else len(rest))
        return end
    raise AssertionError(f"no prose citation matching {pat}")


def insert_figures(body, caps, embed):
    ends = [(first_cite_paragraph_end(body, rf"\bFigures? (?:\d+(?:, | and | to ))*{n}\b"), n) for n in range(1, len(caps) + 1)]
    out, last = [], 0
    for end, n in sorted(ends):
        path, width = embed[n]
        out += [body[last:end], f"\n\n![**Figure {n}.** {caps[n - 1]}]({path.as_posix()}){{width={width:.2f}cm}}\n"]
        last = end
    return "".join(out) + body[last:]


# ---------------------------------------------------------------- template and Word post-processing
def reference_from_template(template, dest):
    """MDPI ships the Fire template as .dot; the file is a .dotx zip. Pandoc and python-docx need the
    document content type, so only [Content_Types].xml changes."""
    assert template and template.exists(), "Fire needs --template <fire-template.dot> (MDPI Word template)"
    zin = zipfile.ZipFile(template)
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
            zout.writestr(item, data)
    ids = {s.style_id for s in Document(dest).styles}
    need = {"MDPI11articletype", "MDPI12title", "MDPI13authornames", "MDPI16affiliation", "MDPI17abstract",
            "MDPI18keywords", "MDPI21heading1", "MDPI22heading2", "MDPI23heading3", "MDPI31text", "MDPI38bullet",
            "MDPI39equation", "MDPI41tablecaption", "MDPI42tablebody", "MDPI51figurecaption", "MDPI52figure",
            "MDPI62backmatter", "MDPI81references"}
    assert need <= ids, f"template lacks styles {sorted(need - ids)}"


def _style(p):
    ppr = p.find(qn("w:pPr"))
    s = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return s.get(qn("w:val")) if s is not None else None


def _set_style(p, sid):
    p.get_or_add_pPr().style = sid


def _ind(p, **kw):
    ppr = p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind"); ppr.append(ind)
    for k, v in kw.items():
        ind.set(qn(f"w:{k}"), str(v))


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"), M_T))


def _border(parent, side, sz):
    b = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "0"), ("color", "auto")):
        b.set(qn(f"w:{k}"), v)
    parent.append(b)


def mdpi_styles(d):
    body = d.element.body
    for p in body.iter(qn("w:p")):
        sid, txt = _style(p), _text(p)
        in_table = any(a.tag == qn("w:tc") for a in p.iterancestors())
        if in_table:
            _set_style(p, "MDPI42tablebody")
            for sz in p.iter(qn("w:sz")):
                sz.set(qn("w:val"), "18")
            for r in p.iter(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
                if rpr.find(qn("w:sz")) is None:
                    e = OxmlElement("w:sz"); e.set(qn("w:val"), "18"); rpr.append(e)
        elif sid in ("Heading1", "Heading2", "Heading3"):
            _set_style(p, {"Heading1": "MDPI21heading1", "Heading2": "MDPI22heading2", "Heading3": "MDPI23heading3"}[sid])
        elif sid in ("CaptionedFigure", "Figure"):
            _set_style(p, "MDPI52figure")
        elif sid == "ImageCaption":
            _set_style(p, "MDPI51figurecaption")
        elif sid == "Bibliography":
            _set_style(p, "MDPI81references")
            ppr = p.get_or_add_pPr()                               # citeproc writes the number as text
            num = OxmlElement("w:numPr"); nid = OxmlElement("w:numId"); nid.set(qn("w:val"), "0")
            num.append(nid); ppr.insert(1, num)
            _ind(p, left=INDENT + 425, hanging=425)
        elif sid is not None and sid.startswith("MDPI"):
            pass                                                   # set through custom-style in the Markdown
        elif p.find(".//" + qn("m:oMathPara")) is not None:
            _set_style(p, "MDPI39equation")
        elif re.match(r"Table \d+\. ", txt):
            _set_style(p, "MDPI41tablecaption")
            if p.find(qn("w:pPr")).find(qn("w:keepNext")) is None:
                p.find(qn("w:pPr")).append(OxmlElement("w:keepNext"))
        elif p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:numPr")) is not None:
            _set_style(p, "MDPI38bullet")                          # the list keeps pandoc's own numbering
            _ind(p, left=INDENT + 425, hanging=283)
        else:
            _set_style(p, "MDPI31text")
    for t in d.tables:
        tblpr = t._tbl.tblPr
        for tag in ("w:tblStyle", "w:tblW", "w:tblInd", "w:tblBorders", "w:jc"):
            for e in tblpr.findall(qn(tag)):
                tblpr.remove(e)
        ncol = len(t.columns)
        width, indent = (TEXT_W, INDENT) if ncol <= 5 else (FULL_W, 0)
        w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(width)); w.set(qn("w:type"), "dxa"); tblpr.append(w)
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(indent)); ind.set(qn("w:type"), "dxa"); tblpr.append(ind)
        bd = OxmlElement("w:tblBorders"); _border(bd, "top", 8); _border(bd, "bottom", 8); tblpr.append(bd)
        need = [max(max(len(_text(c._tc).strip()) for c in col.cells), 4) for col in t.columns]
        widths = [round(width * n / sum(need)) for n in need]
        for gc, cw in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(cw))
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.append(OxmlElement("w:cantSplit"))
            for c, cw in zip(row.cells, widths):
                tcpr = c._tc.get_or_add_tcPr()
                for e in tcpr.findall(qn("w:tcW")) + tcpr.findall(qn("w:tcBorders")):
                    tcpr.remove(e)
                tcw = OxmlElement("w:tcW"); tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(cw)); tcpr.insert(0, tcw)
                if i == 0:
                    tb = OxmlElement("w:tcBorders"); _border(tb, "bottom", 4); tcpr.append(tb)


def word_pdf(docx, pdf):
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx}', $false, $true); $d.ExportAsFixedFormat('{pdf}', 17); "
          f"$d.Close($false); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    assert pdf.exists() and pdf.stat().st_size > 0, pdf


# ---------------------------------------------------------------- checks on the finished .docx
def paragraphs(d):
    return [(p, _style(p) or "", _text(p)) for p in d.element.body.iter(qn("w:p"))]


def check_citations(d):
    pars = paragraphs(d)
    refs = [t for _, s, t in pars if s == "MDPI81references"]
    n_ref = len(refs)
    assert n_ref > 0, "no reference list"
    lead = [int(m.group(1)) if (m := re.match(r"\s*(\d+)\.", t)) else None for t in refs]
    assert lead == list(range(1, n_ref + 1)), f"reference list not numbered 1..{n_ref}: {lead[:5]}..."
    seq = []
    for _, s, t in pars:
        if s == "MDPI81references":
            continue
        assert "???" not in t and not re.search(r"(?<![\w.])@[A-Za-z]", t), f"unresolved citation: {t[:100]}"
        for g in re.findall(r"\[(\d+(?:\s*[–,]\s*\d+)*)\]", t):
            for part in g.split(","):
                a, _, b = part.strip().partition("–")
                seq += list(range(int(a), int(b or a) + 1))
    cited = set(seq)
    assert cited <= set(range(1, n_ref + 1)), f"citations without a reference: {sorted(cited - set(range(1, n_ref + 1)))}"
    assert cited == set(range(1, n_ref + 1)), f"references never cited: {sorted(set(range(1, n_ref + 1)) - cited)}"
    first = list(dict.fromkeys(seq))
    assert first == list(range(1, n_ref + 1)), f"references not numbered in order of first citation: {first[:12]}"
    return n_ref, len(seq)


def _nums(label, t):
    out = set()
    for m in re.finditer(rf"\b{label}s? ((?:\d+(?:, | and | to |–))*\d+)", t):
        parts = re.split(r"(, | and |–| to )", m.group(1))
        nums = [int(x) for x in parts[::2]]
        seps = parts[1::2]
        out.update(nums)
        for i, sep in enumerate(seps):
            if sep in (" to ", "–"):
                out.update(range(nums[i], nums[i + 1] + 1))
    return out


def check_display_items(d, n_fig, n_tab):
    pars = paragraphs(d)
    body_start = next(i for i, (_, s, t) in enumerate(pars) if s == "MDPI21heading1" and t.startswith("1. "))
    body_end = next(i for i, (_, s, t) in enumerate(pars) if s == "MDPI62backmatter")
    prose = {i for i, (p, s, _) in enumerate(pars) if body_start < i < body_end
             and s in ("MDPI31text", "MDPI38bullet", "MDPI39equation")
             and not any(a.tag == qn("w:tc") for a in p.iterancestors())}
    report = {}
    for label, n_items, cap_style in (("Figure", n_fig, "MDPI51figurecaption"), ("Table", n_tab, "MDPI41tablecaption")):
        first = {}
        for i in sorted(prose):
            for n in _nums(label, pars[i][2]):
                first.setdefault(n, i)
        caps = {int(re.match(rf"{label} (\d+)\.", t).group(1)): i for i, (_, s, t) in enumerate(pars)
                if s == cap_style and re.match(rf"{label} \d+\.", t)}
        want = set(range(1, n_items + 1))
        assert set(caps) == want, f"{label} captions {sorted(caps)} != 1..{n_items}"
        assert want <= set(first), f"{label}s not cited in body prose: {sorted(want - set(first))}"
        order = [first[n] for n in sorted(want)]
        assert order == sorted(order), f"{label}s first cited out of numerical order: {order}"
        assert [caps[n] for n in sorted(want)] == sorted(caps.values()), f"{label}s placed out of order"
        late = [n for n in want if caps[n] < first[n]]
        assert not late, f"{label}s placed before their first citation: {late}"
        report[label] = {n: (first[n], caps[n]) for n in sorted(want)}
    return report


def check_numbers(sources, md, d):
    """Decimals: sources -> assembled Markdown (the profile may only add), Markdown <-> .docx (equal)."""
    md = re.sub(r"\{[^{}]*(?:custom-style|width=)[^{}]*\}", "", md)          # markup, not text
    src, asm = Counter(DEC.findall(sources)), Counter(DEC.findall(md))
    lost = src - asm
    assert not lost, f"numbers in the sources that the profile dropped: {dict(lost)}"
    docx = Counter()
    for _, s, t in paragraphs(d):
        if s != "MDPI81references":
            docx.update(DEC.findall(t))
    missing, extra = asm - docx, docx - asm
    assert not missing and not extra, f"Markdown vs .docx decimals: missing {dict(missing)}, extra {dict(extra)}"
    return sum(asm.values())


# ---------------------------------------------------------------- build
def build(template, pdf=True):
    OUT.mkdir(exist_ok=True)
    B.OUT = OUT                                   # B.to_docx writes its temporary reference file here
    FM = B.FM
    kw = FM["keywords"]
    assert 3 <= len(kw) <= 10, "Fire asks for 3 to 10 keywords"

    abstract = re.sub(r"(?m)^# Abstract\s*$", "", B.clean((P / "00_abstract.md").read_text(encoding="utf-8"), "abstract")).strip()
    n_abs = len(abstract.split())
    assert n_abs <= 200, f"abstract has {n_abs} words; Fire asks for about 200 at most"

    raw = {f: B.clean((P / f"{f}.md").read_text(encoding="utf-8"), f) for f in B.BODY}
    EQ = {}
    body = B.equations("\n\n".join(declarations_refs(raw[f], f) for f in B.BODY), EQ)
    unbracketed = re.sub(r"\[[^\[\]]*@[^\[\]]*\]", "", body)
    assert not re.search(r"(?<![\w.])@[A-Za-z]", unbracketed), "bare narrative citation: numeric style would drop the author"
    body = headings(figures_to_words(body))

    decl_src = B.clean((P / f"{B.DECL}.md").read_text(encoding="utf-8"), B.DECL)
    decl = sections(figures_to_words(decl_src))
    assert set(decl) == {"Funding", "Competing interests", "Author contributions", "Data availability",
                         "Code availability", "Ethics approval", "Use of generative AI"}, sorted(decl)
    assert decl["Ethics approval"].startswith("Not applicable."), decl["Ethics approval"][:40]
    funding, coi = funding_and_coi(decl["Funding"], decl["Competing interests"])
    ack = B.ACK.split("\n\n", 1)[1].strip() + " " + genai(decl["Use of generative AI"], raw["03_methods"])

    # figures: 600 dpi files for upload, 300 dpi copies embedded after the first citing paragraph
    caps = [figures_to_words(c) for c in B.captions()]
    tmp = OUT / "_embed"
    tmp.mkdir(exist_ok=True)
    embed, fig_src = {}, {}
    for n in range(1, len(caps) + 1):
        fig_src[n], width = rasterise(n, 600, OUT / f"Figure{n}.png")
        rasterise(n, 300, tmp / f"Figure{n}.png")
        embed[n] = (tmp / f"Figure{n}.png", min(width, 18.4))
    body = insert_figures(body, caps, embed)

    sup_src = declarations_refs(B.clean((P / "supplementary.md").read_text(encoding="utf-8"), "supplementary"),
                                "supplementary")
    sup_list = supplementary_items(sup_src)

    authors = ", ".join(f"{name}^1{',\\*' if i == len(AUTHORS) - 1 else ''}^" for i, (name, _) in enumerate(AUTHORS))
    front = f"""---
title: ""
lang: en-GB
---

::: {{custom-style="MDPI_1.1_article_type"}}
Article
:::

::: {{custom-style="MDPI_1.2_title"}}
{FM['title']}
:::

::: {{custom-style="MDPI_1.3_authornames"}}
\\[AUTHOR ORDER: NEEDS AUTHOR INPUT\\] {authors}
:::

::: {{custom-style="MDPI_1.6_affiliation"}}
^1^ {AFFILIATION}

^\\*^ Correspondence: {CORRESPONDING}
:::

::: {{custom-style="MDPI_1.7_abstract"}}
**Abstract:** {abstract}
:::

::: {{custom-style="MDPI_1.8_keywords"}}
**Keywords:** {'; '.join(kw)}
:::
"""
    das = "\n\n".join([decl["Data availability"], decl["Code availability"]])
    back = f"""::: {{custom-style="MDPI_6.2_back_matter"}}
**Supplementary Materials:** The following supporting information can be downloaded at
https://www.mdpi.com/article/doi/s1: the Supplementary Material (one PDF file), containing {sup_list}.

**Author Contributions:** {author_contributions(decl['Author contributions'])}

**Funding:** {funding}

**Institutional Review Board Statement:** Not applicable.

**Informed Consent Statement:** Not applicable.

**Data Availability Statement:** {das}

**Acknowledgments:** {ack}

**Conflicts of Interest:** {coi}
:::

# References

::: {{#refs}}
:::
"""
    ms_md = front + "\n\n" + body + "\n\n" + back
    (HERE / "manuscript_fire_assembled.md").write_text(ms_md, encoding="utf-8")   # build record, not submitted

    bib = fire_bib(OUT / "_refs.bib")
    order = cite_order(ms_md, bib)
    sup_keys = set(re.findall(r"@([A-Za-z][\w]*)", "".join(re.findall(r"\[[^\[\]]*@[^\[\]]*\]", sup_src))))
    assert sup_keys <= set(order), f"supplement cites keys absent from the main text: {sorted(sup_keys - set(order))}"

    # ---- manuscript.docx in the MDPI template
    ref = OUT / "_template_reference.docx"
    reference_from_template(template, ref)
    out = OUT / "manuscript.docx"
    pypandoc.convert_text(
        ms_md, "docx", format="markdown+tex_math_dollars+pipe_tables+footnotes+superscript",
        outputfile=str(out),
        extra_args=["--citeproc", f"--bibliography={bib}", f"--csl={CSL}",
                    f"--reference-doc={ref}", "--resource-path=" + str(P), "--columns=1000"])
    d = Document(out)
    mdpi_styles(d)
    cp = d.core_properties
    cp.title, cp.author, cp.last_modified_by = FM["title"], "", ""
    cp.comments = cp.keywords = cp.subject = cp.category = cp.identifier = ""
    cp.revision = 1
    d.save(out)
    B.strip_custom_props(out)
    ref.unlink()
    bib.unlink()
    shutil.rmtree(tmp)
    d = Document(out)

    n_ref, n_cit = check_citations(d)
    items = check_display_items(d, len(caps), len(re.findall(r"(?m)^\*\*Table \d+\. ", body)))
    sources = "\n".join([abstract, "\n".join(raw.values()), "\n".join(B.captions()), decl_src])
    n_dec = check_numbers(sources, ms_md, d)

    # ---- supplement.pdf: the shared Supplementary Material, citations numbered as in the main list
    SEQ = {}
    sup = B.equations(figures_to_words(sup_src), SEQ, prefix="S", external=EQ)
    sup = number_citations(sup, order)
    nums = {int(x) for g in re.findall(r"\[(\d+(?:[–,]\d+)*)\]", sup) for x in re.split(r"[–,]", g)}
    assert nums and max(nums) <= n_ref, f"supplement cites numbers outside 1..{n_ref}"
    sup = re.sub(r"\A# Supplementary Material\s*\n", "", sup)
    note = ("Numbers in square brackets refer to the reference list of the main article; the Supplementary "
            "Material has no reference list of its own.")
    sup_md = (f'---\ntitle: "Supplementary Material"\nsubtitle: "{FM["title"]}"\nlang: en-GB\n---\n\n*{note}*\n\n' + sup)
    sup_docx = OUT / "_supplement.docx"
    B.to_docx(sup_md, sup_docx, "Supplementary Material: " + FM["title"])
    if pdf:
        word_pdf(sup_docx, OUT / "supplement.pdf")
    sup_docx.unlink()

    # ---- graphical abstract (drawn and checked by paper/figures/graphical_abstract.py)
    assert GA.exists(), "run paper/figures/graphical_abstract.py first"
    from PIL import Image
    with Image.open(GA) as im:                                        # MDPI: RGB, 8 bit per channel
        assert im.width >= 1100 and im.height >= 560, im.size        # MDPI minimum, w x h
        flat_ga = Image.new("RGB", im.size, "white")
        flat_ga.paste(im.convert("RGBA"), mask=im.convert("RGBA").getchannel("A"))
        flat_ga.save(OUT / "graphical_abstract.png", dpi=im.info.get("dpi", (300, 300)))

    # ---- cover letter (Fire requires one; the two statements are MDPI's wording)
    letter = (HERE / "fire_cover_letter.md").read_text(encoding="utf-8").replace("{{TITLE}}", FM["title"])
    for s in ("We confirm that neither the manuscript nor any parts of its content are currently under consideration "
              "for publication with or published in another journal.",
              "All authors have approved the manuscript and agree with its submission to *Fire*."):
        assert s in flat(letter), f"cover letter lacks the required statement: {s[:50]}"
    stray = set(DEC.findall(letter)) - set(DEC.findall(abstract))
    assert not stray, f"cover letter numbers not in the abstract: {stray}"
    B.to_docx(letter, OUT / "cover_letter.docx", "Cover letter")

    manifest(fig_src, n_abs, n_ref, n_cit, items, n_dec)
    print(f"manuscript.docx: {n_abs}-word abstract, {len(kw)} keywords, {len(caps)} figures embedded, "
          f"{n_ref} references, {n_cit} citation numbers, {n_dec} decimals matched")
    print("display items (first citing paragraph, caption paragraph):", json.dumps(items))
    print("NOTE: the author list carries [AUTHOR ORDER: NEEDS AUTHOR INPUT]")


def manifest(fig_src, n_abs, n_ref, n_cit, items, n_dec):
    commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    dirty = subprocess.run(["git", "status", "--porcelain", "--", "paper", ":!paper/submission_fire"],
                           cwd=ROOT, capture_output=True, text=True).stdout.strip()
    state = commit + (" (with uncommitted source changes)" if dirty else "")
    src = {"manuscript.docx": "paper/0*.md, figure_captions.tex, frontmatter.json, REFERENCES.bib via build_docx.py --journal fire",
           "supplement.pdf": "paper/supplementary.md via build_docx.py --journal fire and Word PDF export",
           "graphical_abstract.png": "paper/figures/graphical_abstract.py, flattened to RGB",
           "cover_letter.docx": "paper/tex/fire_cover_letter.md (draft: date and signature to add)"}
    for n, f in fig_src.items():
        src[f"Figure{n}.png"] = f"{f.relative_to(ROOT).as_posix()}, rasterised at 600 dpi"
    rows = []
    for name in ["manuscript.docx", "supplement.pdf", *[f"Figure{n}.png" for n in fig_src], "graphical_abstract.png",
                 "cover_letter.docx"]:
        f = OUT / name
        if f.exists():
            rows.append(f"| `{name}` | {f.stat().st_size} | `{hashlib.sha256(f.read_bytes()).hexdigest()}` | {src[name]} |")
    text = f"""# Submission package: Fire (MDPI)

Second target; the Natural Hazards package in `paper/submission/` is unchanged. Built from commit
`{state}` by `python paper/tex/build_docx.py --journal fire --template <fire-template.dot>` (the MDPI
template is not in the repository). `manuscript.docx` is built in the MDPI template, which MDPI
licenses for submission only and not for posting online, so it is not committed to this public
repository; rebuild it and compare its SHA-256 below. The text, figures, supplement source and references are the ones
the Natural Hazards build uses; the differences are in `paper/tex/journal_fire.py`.

Build-time checks: abstract {n_abs} words (Fire: about 200 at most); every citation resolves to one of
{n_ref} references, every reference is cited, numbered in order of first citation ({n_cit} citation
numbers); every figure and table cited in body prose, in order, and placed after its first citation;
{n_dec} decimals identical between the assembled Markdown and the .docx. Run separately:
`paper/figures/check_all.py`, `paper/code/check_stale_values.py`, `paper/code/verify_references.py`.

Open items: the author order (`[AUTHOR ORDER: NEEDS AUTHOR INPUT]` on the title page) and every
author's e-mail address, which MDPI publishes; the cover letter's date, signature and the check that
the manuscript is not under consideration elsewhere when it is sent; three suggested reviewers in the
submission system.

| File | Size (bytes) | SHA-256 | Source |
|---|---:|---|---|
""" + "\n".join(rows) + "\n"
    (OUT / "MANIFEST.md").write_text(text, encoding="utf-8")
